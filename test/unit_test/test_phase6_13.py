"""Slide-aware Presentation chunks keep titles, notes, tables and original numbering."""

import hashlib
import json
import pytest
from rag.chunking import factory
from rag.parsers import pptx
from rag.parsers.types import ParseReport, SourceFile

pytestmark = pytest.mark.unit


def parsed(repo_root):
    path = repo_root / "test/fixtures/phase6/presentation-original.pptx"
    gold = json.loads(path.with_name("presentation-original-annotation.json").read_text())
    assert hashlib.sha256(path.read_bytes()).hexdigest() == gold["sha256"]
    report = ParseReport()
    return list(pptx.parse(SourceFile(path, "pptx", path.stat().st_size), report)), report, gold


def test_original_pptx_has_typed_slide_identity_tables_notes_and_blank_gap(repo_root):
    blocks, report, gold = parsed(repo_root)
    assert all(getattr(b, "slide", None) is not None for b in blocks), "Typed original slide identity is missing"
    assert sorted({b.slide.number for b in blocks}) == gold["nonempty_slide_numbers"]
    assert report.pages_total == 4 and report.pages_without_text == 1
    assert [b.table.cells[2].text for b in blocks if b.table] == ["Height"]
    assert all(fact in "\n".join(b.text for b in blocks) for fact in gold["facts"])


def test_presentation_chunks_repeat_title_once_never_mix_slides_and_bound_notes(repo_root):
    blocks, _, gold = parsed(repo_root)
    assert factory.available("presentation"), "Presentation implementation is missing"
    chunks = list(factory.get_chunker("presentation")(blocks, factory.get_chunker_config("presentation", {"chunk_token_num": 64})))
    assert sorted({c.slide.number for c in chunks}) == gold["nonempty_slide_numbers"]
    assert all(c.text.count(c.slide.title) == 1 and c.token_num <= 64 for c in chunks)
    assert len([c for c in chunks if c.slide.number == 4]) > 1
    assert len([c for c in chunks if c.slide.number == 3]) == 1
    assert all("ORCHARD" not in c.text and "O9100" not in c.text for c in chunks if c.slide.number == 1)
    assert all("LIGHTHOUSE" not in c.text and "73 units" not in c.text for c in chunks if c.slide.number == 4)
    assert all(any(fact in c.text for c in chunks) for fact in gold["facts"])
    assert any(c.tables and c.tables[0].cells[3].text == "52" for c in chunks if c.slide.number == 1)
    assert [c.index for c in chunks] == list(range(len(chunks)))


def test_slide_index_and_public_view_have_identity_without_fabricated_boxes():
    from rag.svr.chunk_doc import build_row
    from rag.chunking.types import Chunk
    from api.db.services.chunk_view import chunk_view_fields
    from rag.parsers.slides import Slide

    chunk = Chunk("Closing title only", (), {}, 4, 0, slide=Slide(3, "Closing title only"))
    row = build_row(chunk, index=0, doc_id="a" * 32, kb_id="b" * 32, doc_name="slides.pptx", task_id="c" * 32, from_page=0, to_page=1000000, vector=[1.0], enabled=True)
    assert chunk_view_fields(row).get("slide") == {"number": 3, "title": "Closing title only"}
    assert row["position_int"] == [] and row["page_dims_flt"] == []


def test_missing_or_reordered_slide_identity_is_refused():
    from rag.parsers.types import ParsedBlock
    from rag.parsers.slides import Slide
    from rag.parsers.errors import FileDamaged

    run = lambda blocks: list(factory.get_chunker("presentation")(blocks, factory.get_chunker_config("presentation", {})))
    with pytest.raises(FileDamaged):
        run([ParsedBlock("unknown slide")])
    with pytest.raises(FileDamaged):
        run([ParsedBlock("third", slide=Slide(3)), ParsedBlock("first", slide=Slide(1))])
    assert run([]) == []


def test_small_budget_refuses_title_and_long_notes_keep_images():
    from rag.parsers.types import ParsedBlock
    from rag.parsers.slides import Slide
    from rag.chunking.errors import ChunkTooLong

    block = ParsedBlock("long notes " * 100, slide=Slide(2, "Slide title"), image_ids=("a" * 32,))
    run = lambda size: list(factory.get_chunker("presentation")([block], factory.get_chunker_config("presentation", {"chunk_token_num": size})))
    with pytest.raises(ChunkTooLong):
        run(1)
    chunks = run(30)
    assert len(chunks) > 1 and all(c.image_ids == ("a" * 32,) and c.slide.number == 2 for c in chunks)


def test_pptx_respects_text_limit(repo_root):
    from rag.parsers.types import ParseLimits
    from rag.parsers.errors import FileTooLarge

    path = repo_root / "test/fixtures/phase6/presentation-original.pptx"
    with pytest.raises(FileTooLarge):
        list(pptx.parse(SourceFile(path, "pptx", path.stat().st_size), ParseReport(), limits=ParseLimits(max_text_bytes=32)))


def test_office_table_public_decoder_accepts_safe_cells_without_invented_crop():
    from rag.parsers.tables import Table,TableCell
    from rag.parsers.table_features import table_views
    table=Table(1,1,(TableCell(0,0,'<script>52</script>'),))
    assert table_views([table.to_dto()])==[table.to_dto()]
    assert '&lt;script&gt;' in table.html()


def test_merged_slide_cells_and_title_only_keep_original_number(tmp_path):
    import pptx as library
    from pptx.util import Inches
    deck=library.Presentation()
    first=deck.slides.add_slide(deck.slide_layouts[1]);first.shapes.title.text='Merged title'
    table=first.shapes.add_table(2,2,Inches(1),Inches(1),Inches(5),Inches(2)).table
    table.cell(0,0).merge(table.cell(0,1));table.cell(0,0).text='Merged header'
    table.cell(1,0).text='height';table.cell(1,1).text='52'
    last=deck.slides.add_slide(deck.slide_layouts[1]);last.shapes.title.text='Only title'
    path=tmp_path/'merged.pptx';deck.save(path)
    blocks=list(pptx.parse(SourceFile(path,'pptx',path.stat().st_size),ParseReport()))
    grid=next(b.table for b in blocks if b.table)
    assert grid.cells[0].colspan==2 and len(grid.cells)==3
    chunks=list(factory.get_chunker('presentation')(blocks,factory.get_chunker_config('presentation',{})))
    assert [c.slide.number for c in chunks]==[1,2]
    assert chunks[1].text.strip()=='Only title'
