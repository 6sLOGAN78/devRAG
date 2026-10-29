"""PPTX parser (PARSE-11): typed original slide groups, configured remote figures; legacy skips counted (D-27)."""

from __future__ import annotations

import zipfile
from collections.abc import Iterator

import pptx
from lxml import etree
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError

from common.settings import Settings
from rag.parsers.errors import FileDamaged, NoReadableText, FileTooLarge
from rag.parsers.slides import Slide
from rag.parsers.tables import Table, TableCell
from rag.parsers.limits import check_ooxml_limits
from rag.parsers.ooxml_images import OOXMLImages
from rag.parsers.types import ParsedBlock, ParseLimits, ParseReport, SourceFile

_DAMAGE = (zipfile.BadZipFile, KeyError, ValueError, AttributeError, PackageNotFoundError, etree.XMLSyntaxError, OSError)


def _shape_blocks(shapes, skip_id, identity):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _shape_blocks(shape.shapes, skip_id, identity)
        elif getattr(shape, "has_table", False):
            grid = shape.table
            if len(grid.rows) > 256 or len(grid.columns) > 128 or len(grid.rows) * len(grid.columns) > 4096:
                raise FileTooLarge()
            cells = tuple(TableCell(r, c, cell.text.strip(), cell.span_height, cell.span_width, r == 0) for r, row in enumerate(grid.rows) for c, cell in enumerate(row.cells) if not cell.is_spanned)
            table = Table(len(grid.rows), len(grid.columns), cells)
            yield ParsedBlock(table.plain_text(), "table", table=table, slide=identity)
        elif getattr(shape, "has_text_frame", False) and shape.shape_id != skip_id:
            text = shape.text_frame.text.strip()
            if text:
                yield ParsedBlock(text, "slide", slide=identity)


def _extract(source: SourceFile, report: ParseReport, settings: Settings | None = None, limits: ParseLimits | None = None) -> list[ParsedBlock]:
    prs = pptx.Presentation(str(source.path))
    limits = limits or ParseLimits()
    if len(prs.slides) > min(limits.max_pages, 20000):
        raise FileTooLarge()
    report.pages_total = len(prs.slides)
    blocks = []
    total_text = 0
    with OOXMLImages(source.path, settings) as collector:
        for number, slide in enumerate(prs.slides, 1):
            title = slide.shapes.title
            label = title.text_frame.text.strip() if title is not None else ""
            if len(label) > 1024:
                raise FileTooLarge()
            identity = Slide(number, label)
            group = [ParsedBlock(label, "slide", slide=identity)] if label else []
            group.extend(_shape_blocks(slide.shapes, title.shape_id if title is not None else None, identity))
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes is not None and notes.text.strip():
                    group.append(ParsedBlock(notes.text.strip(), "slide", slide=identity))
            assets = collector.for_element(slide._element, str(slide.part.partname))
            if assets:
                group.append(ParsedBlock("", "slide", assets=assets, slide=identity))
            if not any(b.text.strip() for b in group):
                report.pages_without_text += 1
            total_text += sum(len(b.text.encode()) for b in group)
            if total_text > limits.max_text_bytes:
                raise FileTooLarge()
            blocks.extend(group)
        if not collector.enabled:
            # Retain the legacy report for pictures when remote normalization is disabled.
            def pictures(shapes):
                return sum(pictures(s.shapes) if s.shape_type == MSO_SHAPE_TYPE.GROUP else int(s.shape_type == MSO_SHAPE_TYPE.PICTURE) for s in shapes)

            report.images_skipped = sum(pictures(slide.shapes) for slide in prs.slides)
            if report.images_skipped:
                report.add_note(f"{report.images_skipped} embedded images were skipped")
    return blocks


def parse(
    source: SourceFile,
    report: ParseReport,
    *,
    page_range: tuple[int, int] | None = None,
    limits: ParseLimits | None = None,
    settings: Settings | None = None,
) -> Iterator[ParsedBlock]:
    del page_range
    check_ooxml_limits(source.path)
    try:
        blocks = _extract(source, report, settings, limits)
    except _DAMAGE:
        raise FileDamaged() from None
    if not blocks:
        raise NoReadableText()
    return iter(blocks)
