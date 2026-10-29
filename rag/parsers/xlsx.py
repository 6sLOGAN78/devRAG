"""XLSX parser (PARSE-10): read-only, cached values, header row repeated per group, cell cap."""

from __future__ import annotations

import zipfile
from collections.abc import Iterator

import openpyxl
from lxml import etree

from rag.parsers.errors import FileDamaged, FileTooLarge, NoReadableText
from rag.parsers.limits import check_ooxml_limits
from rag.parsers.tabletext import render_table
from rag.parsers.types import ParsedBlock, ParseLimits, ParseReport, SourceFile

GROUP_ROWS = 20
_DAMAGE = (zipfile.BadZipFile, KeyError, ValueError, AttributeError, etree.XMLSyntaxError, OSError)


def _extract(source: SourceFile, report: ParseReport, max_cells: int) -> list[ParsedBlock]:
    wb = openpyxl.load_workbook(str(source.path), read_only=True, data_only=True)
    blocks: list[ParsedBlock] = []
    cells = 0
    stopped = False
    try:
        for ws in wb.worksheets:
            header: list[object] | None = None
            group: list[list[object]] = []

            def flush() -> None:
                if group:
                    rows = [header, *group] if header is not None else list(group)
                    blocks.append(ParsedBlock(f"{ws.title}\n{render_table(rows)}", "sheet"))
                    group.clear()

            for row in ws.iter_rows(values_only=True):
                if not any(c is not None and str(c).strip() for c in row):
                    continue
                cells += len(row)
                if cells > max_cells:
                    stopped = True
                    break
                if header is None:
                    header = list(row)
                    continue
                group.append(list(row))
                if len(group) >= GROUP_ROWS:
                    flush()
            flush()
            if header is not None and not blocks_for(blocks, ws.title):
                blocks.append(ParsedBlock(f"{ws.title}\n{render_table([header])}", "sheet"))
            if stopped:
                break
    finally:
        wb.close()
    if stopped:
        report.add_note(f"reading stopped after {max_cells} cells")
    return blocks


def blocks_for(blocks: list[ParsedBlock], title: str) -> bool:
    return any(b.text.startswith(f"{title}\n") for b in blocks)


def _extract_records(source: SourceFile, max_cells: int) -> list[ParsedBlock]:
    from rag.parsers.records import MAX_RECORDS, row_record, pointer_part

    from datetime import date, datetime, time

    wb = openpyxl.load_workbook(str(source.path), read_only=True, data_only=True)
    blocks = []
    cells = 0
    try:
        for ws in wb.worksheets:
            header = None
            for number, row in enumerate(ws.iter_rows(values_only=True), 1):
                values = list(row)
                while values and values[-1] in (None, ""):
                    values.pop()
                if not values:
                    continue
                cells += len(row)
                if cells > max_cells or len(blocks) >= MAX_RECORDS:
                    raise FileTooLarge()
                prefix = "/" + pointer_part(ws.title) + "/" + str(number)
                if header is None:
                    header = values
                    row_record(header, [""] * len(header), prefix)
                    continue
                if len(values) > len(header):
                    raise FileDamaged()
                values += [None] * (len(header) - len(values))
                values = [value.isoformat() if isinstance(value, (datetime, date, time)) else value for value in values]
                record = row_record(header, values, prefix, (number, number))
                blocks.append(ParsedBlock(record.text(), "sheet", record=record, line_range=record.line_range))
    finally:
        wb.close()
    return blocks


def _extract_tables(source, max_cells):
    from rag.parsers.tables import Table, TableCell

    # Stream the actual XML cells before materializing merge information. Ignore
    # untrusted worksheet dimensions so a forged small range cannot bypass the cap.
    probe = openpyxl.load_workbook(str(source.path), read_only=True, data_only=True)
    seen = 0
    try:
        for sheet in probe.worksheets:
            sheet.reset_dimensions()
            for row in sheet.iter_rows(values_only=True):
                seen += len(row)
                if seen > max_cells:
                    raise FileTooLarge()
    finally:
        probe.close()
    wb = openpyxl.load_workbook(str(source.path), read_only=False, data_only=True)
    blocks = []
    total = 0
    try:
        for ws in wb.worksheets:
            rows, columns = ws.max_row, ws.max_column
            if not any(cell.value is not None for row in ws.iter_rows(min_row=1, max_row=1, max_col=min(columns, 129)) for cell in row):
                continue
            total += rows * columns
            if total > max_cells or columns > 128:
                raise FileTooLarge()
            merges = list(ws.merged_cells.ranges)
            depth = max([1, *[m.max_row for m in merges if m.min_row == 1]])
            if any(m.min_row == 1 and m.max_col > m.min_col for m in merges) and rows > 1 and all(isinstance(ws.cell(2, c).value, str) for c in range(1, columns + 1)):
                depth = max(depth, 2)
            if depth >= rows:
                continue
            data_size = min(256, 4096 // columns) - depth
            if data_size < 1:
                raise FileTooLarge()
            for start in range(depth + 1, rows + 1, data_size):
                stop = min(rows, start + data_size - 1)
                physical = list(range(1, depth + 1)) + list(range(start, stop + 1))
                selected = {r: i for i, r in enumerate(physical)}
                cells = []
                occupied = set()
                for r in physical:
                    for c in range(1, columns + 1):
                        if (r, c) in occupied:
                            continue
                        merge = next((m for m in merges if m.min_row <= r <= m.max_row and m.min_col <= c <= m.max_col), None)
                        end = min(merge.max_row, stop if r >= start else depth) if merge else r
                        width = merge.max_col - c + 1 if merge else 1
                        value = ws.cell(merge.min_row, merge.min_col).value if merge else ws.cell(r, c).value
                        cells.append(TableCell(selected[r], c - 1, "" if value is None else str(value), end - r + 1, width, r <= depth))
                        occupied.update((rr, cc) for rr in range(r, end + 1) for cc in range(c, c + width))
                table = Table(len(physical), columns, tuple(cells), ws.title)
                blocks.append(ParsedBlock(table.plain_text(), "table", table=table, line_range=(start, stop)))
    finally:
        wb.close()
    return blocks


def parse(
    source: SourceFile,
    report: ParseReport,
    *,
    page_range: tuple[int, int] | None = None,
    limits: ParseLimits | None = None,
    structured_records: bool = False,
    structured_tables: bool = False,
) -> Iterator[ParsedBlock]:
    del page_range
    limits = limits or ParseLimits()
    check_ooxml_limits(source.path)
    try:
        blocks = (
            _extract_tables(source, limits.xlsx_max_cells)
            if structured_tables
            else (_extract_records(source, limits.xlsx_max_cells) if structured_records else _extract(source, report, limits.xlsx_max_cells))
        )
    except _DAMAGE:
        raise FileDamaged() from None
    if not blocks:
        raise NoReadableText()
    return iter(blocks)
