from pathlib import Path
from pptx import Presentation

from .base import BaseParser
from ..models import DocumentStructure, TextBlock, TableBlock, PageNode
from ..errors import ParsingFailureError

class PptParser(BaseParser):
    def parse(self, file_path: Path) -> DocumentStructure:
        self._validate_file(file_path)
        
        try:
            prs = Presentation(file_path)
        except Exception as e:
            raise ParsingFailureError(f"Failed to read PPTX file: {e}")

        pages = []
        for i, slide in enumerate(prs.slides):
            blocks = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    blocks.append(TextBlock(text=shape.text.strip(), block_id=f"slide_{i}_shape_{shape.shape_id}"))
                if shape.has_table:
                    rows_data = []
                    max_cols = 0
                    for row in shape.table.rows:
                        row_text = [cell.text.strip() for cell in row.cells]
                        rows_data.append(row_text)
                        if len(row_text) > max_cols:
                            max_cols = len(row_text)
                    if rows_data:
                        blocks.append(TableBlock(rows=rows_data, columns=max_cols, block_id=f"slide_{i}_table_{shape.shape_id}"))
            
            pages.append(PageNode(page_number=i+1, blocks=blocks))
        
        return DocumentStructure(
            source_path=str(file_path),
            file_type="pptx",
            pages=pages,
            metadata={"total_slides": len(pages)}
        )
