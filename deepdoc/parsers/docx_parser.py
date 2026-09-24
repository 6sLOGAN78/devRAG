from pathlib import Path
from docx import Document

from .base import BaseParser
from ..models import DocumentStructure, TextBlock, TableBlock, PageNode
from ..errors import ParsingFailureError

class DocxParser(BaseParser):
    def parse(self, file_path: Path) -> DocumentStructure:
        self._validate_file(file_path)
        
        try:
            doc = Document(file_path)
        except Exception as e:
            raise ParsingFailureError(f"Failed to read DOCX file: {e}")

        blocks = []
        for i, para in enumerate(doc.paragraphs):
            text = para.text.strip()
            if text:
                blocks.append(TextBlock(text=text, block_id=f"p_{i}"))

        for i, table in enumerate(doc.tables):
            rows_data = []
            max_cols = 0
            for row in table.rows:
                row_text = [cell.text.strip() for cell in row.cells]
                rows_data.append(row_text)
                if len(row_text) > max_cols:
                    max_cols = len(row_text)
            if rows_data:
                blocks.append(TableBlock(rows=rows_data, columns=max_cols, block_id=f"t_{i}"))
                
        # Word documents don't have explicit pages in `python-docx`. We put everything in page 1.
        page = PageNode(page_number=1, blocks=blocks)
        
        return DocumentStructure(
            source_path=str(file_path),
            file_type="docx",
            pages=[page],
            metadata={"total_blocks": len(blocks)}
        )
